#!/usr/bin/env python3
"""Bounded, non-executing file parser. Reads bytes from stdin; emits one JSON object.

This is a parser, not a malware scanner. The production worker calls it only AFTER
ClamAV accepts the original object. No OCR, macros, formulas or links are executed.
OS memory/CPU limits are additional safeguards, not a complete sandbox.
"""
from __future__ import annotations
import base64, csv, datetime as dt, io, json, os, re, subprocess, sys, tempfile, zipfile
from pathlib import PurePosixPath

MAX_BYTES = 10 * 1024 * 1024
MAX_XML_BYTES = 40 * 1024 * 1024
MAX_ROWS, MAX_COLS, MAX_CELLS, MAX_SHEETS = 2000, 80, 60000, 10
MAX_TEXT, MAX_PAGES = 2_000_000, 30

class InputError(Exception):
    def __init__(self, code: str, message: str):
        self.code, self.message = code, message
        super().__init__(message)

def reject(code: str, message: str):
    raise InputError(code, message)

def limits():
    # In the Linux container these limits also apply to pdfinfo/pdftotext children.
    try:
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024**2, 768 * 1024**2))
        resource.setrlimit(resource.RLIMIT_CPU, (20, 22))
        resource.setrlimit(resource.RLIMIT_FSIZE, (8 * 1024**2, 8 * 1024**2))
        resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    except (ImportError, ValueError, OSError):
        pass  # Windows: outer Node wall/output limits remain; use Docker for isolation.

def label(index: int) -> str:
    text = ''
    while index:
        index, tail = divmod(index - 1, 26)
        text = chr(65 + tail) + text
    return text

def clean_value(value):
    if value is None: return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    if isinstance(value, bool): return 'TRUE' if value else 'FALSE'
    text = str(value)
    if len(text) > 2000: reject('CELL_TOO_LONG', 'A cell exceeds 2000 characters. Split narrative content from the data table.')
    return text

def inspect_zip(data: bytes):
    from defusedxml import ElementTree as ET
    try: z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile: reject('INVALID_XLSX', 'This is not a readable XLSX package.')
    infos = z.infolist()
    if len(infos) > 1500: reject('WORKBOOK_LIMIT', 'Workbook has too many internal parts.')
    if len({i.filename for i in infos}) != len(infos): reject('DUPLICATE_ZIP_PART', 'Duplicate workbook parts are not accepted.')
    if sum(i.file_size for i in infos) > MAX_XML_BYTES: reject('EXPANDED_LIMIT', 'Expanded workbook exceeds 40 MiB.')
    names = {i.filename for i in infos}
    if not {'[Content_Types].xml', 'xl/workbook.xml'}.issubset(names): reject('INVALID_XLSX', 'The ZIP is not an XLSX workbook.')
    merges, hidden_rows = {}, {}
    for info in infos:
        name, lower = info.filename, info.filename.lower()
        if '\\' in name or name.startswith('/') or '..' in PurePosixPath(name).parts:
            reject('UNSAFE_ZIP_PART', 'Unsafe workbook part name.')
        if info.flag_bits & 1: reject('ENCRYPTED_WORKBOOK', 'Unlock the workbook before uploading.')
        if info.file_size > 12 * 1024**2: reject('PART_LIMIT', 'A workbook part is too large.')
        if info.file_size > 1_000_000 and info.file_size / max(info.compress_size, 1) > 200:
            reject('ZIP_RATIO', 'Workbook compression ratio exceeds the safe parser limit.')
        if any(token in lower for token in ('vbaproject', 'externallinks/', 'embeddings/', 'activex/', 'connections.xml')):
            reject('ACTIVE_WORKBOOK_CONTENT', 'Macros, embedded objects and external workbook data are not accepted. Save a values-only XLSX.')
        if name.endswith(('.xml', '.rels')):
            content = z.read(info)
            if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                reject('XML_ENTITY', 'XML document types/entities are not accepted.')
            try: root = ET.fromstring(content)
            except Exception: reject('INVALID_XML', 'A workbook XML part is invalid.')
            if name.endswith('.rels') and any(n.attrib.get('TargetMode') == 'External' for n in root.iter()):
                reject('EXTERNAL_LINK', 'External workbook relationships are not accepted. Remove links before upload.')
            if name.startswith('xl/worksheets/'):
                merges[name] = [n.attrib['ref'] for n in root.iter() if n.tag.endswith('}mergeCell')]
                hidden_rows[name] = [int(n.attrib['r']) for n in root.iter() if n.tag.endswith('}row') and n.attrib.get('hidden') in ('1', 'true')]
    return z, merges, hidden_rows

def xlsx(data: bytes):
    z, merged_by_part, hidden_by_part = inspect_zip(data)
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    if len(wb.worksheets) > MAX_SHEETS: reject('SHEET_LIMIT', 'At most 10 sheets per workbook are supported.')
    sheets, count, chars = [], 0, 0
    for ws in wb.worksheets:
        part = getattr(ws, '_worksheet_path', '')
        merged = merged_by_part.get(part, [])
        hidden = hidden_by_part.get(part, [])
        merge_lookup = {}
        for ref in merged:
            a, b, c, d = range_boundaries(ref)
            if c > MAX_COLS or d > MAX_ROWS or (c-a+1)*(d-b+1) > MAX_CELLS:
                reject('MERGED_RANGE_LIMIT', 'Merged range exceeds the supported table size.')
            anchor = f'{label(a)}{b}'
            for r in range(b, d+1):
                for col in range(a, c+1):
                    if r != b or col != a: merge_lookup[(r,col)] = anchor
        ws.reset_dimensions()  # Do not trust a forged or stale worksheet dimension.
        rows = []
        for rn, row in enumerate(ws.iter_rows(), 1):
            if rn > MAX_ROWS: reject('ROW_LIMIT', 'A sheet exceeds 2000 physical rows. Split the workbook.')
            if len(row) > MAX_COLS: reject('COLUMN_LIMIT', 'A sheet exceeds 80 columns. Remove unused formatted columns.')
            cells = []
            for cn in range(1,max(len(row),max((col for (r,col) in merge_lookup if r==rn),default=0))+1):
                cell = row[cn-1] if cn<=len(row) else None
                v = cell.value if cell is not None else None
                if v is None and (rn,cn) not in merge_lookup: continue
                count += 1
                if count > MAX_CELLS: reject('CELL_LIMIT', 'Workbook exceeds 60000 populated/merged cells.')
                value = clean_value(v)
                chars += len(value or '')
                if chars > MAX_TEXT: reject('TEXT_LIMIT', 'Workbook text exceeds 2 million characters.')
                typ = 'formula' if getattr(cell,'data_type',None) == 'f' else ('date' if isinstance(v,(dt.date,dt.datetime)) else ('number' if isinstance(v,(int,float)) and not isinstance(v,bool) else ('error' if getattr(cell,'data_type',None)=='e' else 'string')))
                item = {'column':label(cn),'address':f'{label(cn)}{rn}','value':None if typ=='formula' else value,'type':typ}
                if typ == 'formula': item['formula'] = value
                if (rn,cn) in merge_lookup: item['mergedFrom'] = merge_lookup[(rn,cn)]
                cells.append(item)
            if cells: rows.append({'index':rn,'hidden':rn in hidden,'cells':cells})
        sheets.append({'name':ws.title,'hidden':ws.sheet_state!='visible','rows':rows,'mergedRanges':merged})
    wb.close(); z.close()
    return {'kind':'WORKBOOK','method':'OPENPYXL_VALUES_NO_FORMULA_EXECUTION','sheets':sheets,
            'needsHumanReview':True,'warnings':['Formulas are preserved as formulas, never executed or used as consumption. Hidden content is excluded unless explicitly selected.']}

def csv_file(data: bytes):
    try: text = data.decode('utf-8-sig', errors='strict')
    except UnicodeDecodeError: reject('CSV_ENCODING', 'Save CSV as UTF-8 before upload.')
    if '\x00' in text: reject('CSV_BINARY', 'CSV contains binary NUL bytes.')
    if len(text) > MAX_TEXT: reject('TEXT_LIMIT', 'CSV text is too large.')
    try: dialect = csv.Sniffer().sniff(text[:12000], delimiters=',;\t')
    except csv.Error: dialect = csv.excel
    rows=[]; count=0
    for rn,row in enumerate(csv.reader(io.StringIO(text), dialect),1):
        if rn>MAX_ROWS or len(row)>MAX_COLS: reject('CSV_LIMIT','CSV exceeds 2000 rows or 80 columns.')
        cells=[]
        for cn,v in enumerate(row,1):
            if v=='':continue
            count+=1
            if count>MAX_CELLS: reject('CELL_LIMIT','CSV exceeds 60000 cells.')
            clean_value(v)
            # Leading formula text remains untrusted literal input. Never execute it.
            cells.append({'column':label(cn),'address':f'{label(cn)}{rn}','value':v,'type':'string'})
        if cells:rows.append({'index':rn,'hidden':False,'cells':cells})
    return {'kind':'WORKBOOK','method':'UTF8_CSV','sheets':[{'name':'CSV','hidden':False,'rows':rows,'mergedRanges':[]}],
            'needsHumanReview':True,'warnings':['Confirm the detected delimiter and date/number convention before import.']}

def pdf(data: bytes):
    if not data.startswith(b'%PDF-'): reject('INVALID_PDF','PDF signature missing.')
    with tempfile.TemporaryDirectory(prefix='cs-parser-') as directory:
        path=os.path.join(directory,'input.pdf')
        with open(path,'wb') as f:f.write(data)
        try:
            info=subprocess.run(['pdfinfo',path],capture_output=True,timeout=8,check=False)
        except FileNotFoundError: reject('PDF_TOOL_UNAVAILABLE','Install Poppler utilities or use the supplied Docker image.')
        if info.returncode: reject('PDF_UNREADABLE','PDF is damaged, protected, or unsupported. Upload an unlocked PDF or review manually.')
        metadata=info.stdout.decode('utf-8',errors='replace')
        if re.search(r'^Encrypted:\s+yes',metadata,re.M):reject('PDF_ENCRYPTED','Upload an unlocked copy; this parser does not bypass document protection.')
        m=re.search(r'^Pages:\s+(\d+)',metadata,re.M)
        if not m:reject('PDF_UNREADABLE','PDF page count could not be read.')
        count=int(m.group(1))
        if count<1 or count>MAX_PAGES:reject('PAGE_LIMIT','Split PDFs larger than 30 pages before upload.')
        out=os.path.join(directory,'text.txt')
        result=subprocess.run(['pdftotext','-layout','-enc','UTF-8','-f','1','-l',str(count),path,out],capture_output=True,timeout=12,check=False)
        if result.returncode:reject('PDF_EXTRACTION_FAILED','PDF text extraction failed. The original remains available for manual review after scanning.')
        with open(out,'rb') as f:raw=f.read(MAX_TEXT+1)
        if len(raw)>MAX_TEXT:reject('TEXT_LIMIT','PDF extracted text exceeds the safe limit.')
        pages=raw.decode('utf-8',errors='replace').split('\f')
        if pages and not pages[-1].strip():pages.pop()
        pages=pages[:count]
        while len(pages)<count:pages.append('')
        return {'kind':'INVOICE','method':'POPPLER_TEXT','pageCount':count,
                'pages':[{'page':n,'text':text,'status':'TEXT' if len(text.strip())>=20 else 'MANUAL_REQUIRED'} for n,text in enumerate(pages,1)],
                'needsHumanReview':True,'ocrAvailable':False,
                'warnings':['Page text and invoice boundaries are suggestions. Check the original. Image-only pages need manual fields; OCR is not enabled.']}

def xlsx_export(data: bytes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    payload=json.loads(data)
    rows=payload.get('rows',[])
    if len(rows)>1000:reject('EXPORT_LIMIT','Export supports 1000 staged rows per batch.')
    wb=Workbook(); ws=wb.active;ws.title='Normalized data';ws.sheet_view.showGridLines=False
    headers=['Row ID','File','Sheet / Pages','Source row','Status','Target','Quantity','Unit','Start','End','Source / Task','Invoice number','Amount INR','Issues','Review reason','Record ID']
    ws.append(headers)
    for r in rows:
        n=r.get('normalized',{}); src=r.get('source_ref',{})
        values=[r.get('id',''),r.get('filename',''),src.get('sheet') or f"{src.get('pageStart','')} - {src.get('pageEnd','')}",str(src.get('row','')),r.get('status',''),r.get('target',''),n.get('quantity',''),n.get('unit',''),n.get('intervalStart',''),n.get('intervalEnd',''),n.get('sourceCode') or n.get('taskId') or n.get('kpiCode',''),n.get('invoiceNumber',''),n.get('amountInr',''),'; '.join(i.get('code','')+': '+i.get('message','') for i in r.get('issues',[])),r.get('review_reason',''),r.get('record_id') or '']
        ws.append([str(v or '') for v in values])
        for cell in ws[ws.max_row]:
            cell.data_type='s'  # Literal strings, never executable spreadsheet formulas.
            cell.font=Font(name='Calibri',size=11,color='008060')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]:cell.font=Font(name='Calibri',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='173D36');cell.alignment=Alignment(wrap_text=True)
    widths=[38,30,25,12,16,14,17,13,16,16,30,24,18,70,50,38]
    for i,w in enumerate(widths,1):ws.column_dimensions[label(i)].width=w
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
    notes=wb.create_sheet('Read me');notes.column_dimensions['A'].width=110
    for message in ['CarbonSynq normalized staging export','This is a review/export of uploaded data, not a certified emissions report.','Quantities and money are separate. All cell values are literal text to prevent spreadsheet formula injection.','DRAFT / READY is not an approved inventory calculation. Original files, raw fields and cell/page provenance remain in the application.','Source: university uploads and explicitly recorded reviewer corrections.']:
        notes.append([message]);notes.cell(notes.max_row,1).alignment=Alignment(wrap_text=True);notes.row_dimensions[notes.max_row].height=34
    output=io.BytesIO();wb.save(output)
    return {'base64':base64.b64encode(output.getvalue()).decode('ascii')}

def main():
    limits()
    kind=sys.argv[1] if len(sys.argv)>1 else ''
    data=sys.stdin.buffer.read(MAX_BYTES+1)
    if len(data)>MAX_BYTES:reject('FILE_TOO_LARGE','Parser input exceeds 10 MiB.')
    if not data:reject('EMPTY_FILE','Empty file.')
    fn={'xlsx':xlsx,'csv':csv_file,'pdf':pdf,'export':xlsx_export}.get(kind)
    if not fn:reject('UNSUPPORTED_FORMAT','Supported parser formats: XLSX, UTF-8 CSV, PDF.')
    return fn(data)

if __name__=='__main__':
    try: result={'ok':True,'data':main()}
    except InputError as e:result={'ok':False,'error':{'code':e.code,'message':e.message}}
    except subprocess.TimeoutExpired:result={'ok':False,'error':{'code':'PARSER_TIMEOUT','message':'Parsing timed out. Split the file or use reviewed manual input.'}}
    except Exception:result={'ok':False,'error':{'code':'PARSER_FAILED','message':'File could not be parsed safely. Save a values-only XLSX or unlocked PDF, or use manual review.'}}
    sys.stdout.write(json.dumps(result,ensure_ascii=True,separators=(',',':')))
