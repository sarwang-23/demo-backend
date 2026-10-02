"""Build inert parser-rejection fixtures in a caller-supplied temporary directory."""
from pathlib import Path
import datetime, sys, zipfile, io
from openpyxl import Workbook
out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
w=Workbook();s=w.active;s.append(['Consumption','Unit','Date']);s.append([10,'kWh',datetime.date(2026,5,3)]);s.append(['=1+1','kWh',datetime.date(2026,5,4)]);s.row_dimensions[3].hidden=True;s.merge_cells('D1:F1');s['D1']='Merged title';s['C2'].number_format='yyyy-mm-dd';w.save(out/'dates-formula.xlsx')
source=(out/'dates-formula.xlsx').read_bytes()
def altered(name,add=None,replace=None):
 with zipfile.ZipFile(io.BytesIO(source)) as z,zipfile.ZipFile(out/name,'w',zipfile.ZIP_DEFLATED) as o:
  for item in z.infolist():
   val=z.read(item.filename)
   if replace:val=replace(item.filename,val)
   o.writestr(item.filename,val)
  if add:o.writestr(*add)
altered('macro.xlsx',('xl/vbaProject.bin',b'INERT_TEST_NOT_EXECUTABLE'))
altered('traversal.xlsx',('../test.txt',b'INERT_TEST'))
altered('entities.xlsx',replace=lambda n,b: b'<!DOCTYPE root [<!ENTITY x "test">]>'+b if n=='xl/workbook.xml' else b)
altered('external.xlsx',('xl/_rels/test.rels',b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://example.invalid/test" Target="http://example.invalid/never-fetch" TargetMode="External"/></Relationships>'))
altered('zipbomb.xlsx',('huge.txt',b'0'*2_000_000))
w=Workbook();w.active.cell(2001,1,'outside row cap');w.save(out/'rows.xlsx')
w=Workbook();w.active.cell(1,81,'outside column cap');w.save(out/'columns.xlsx')
def simple_pdf(path,pages=31,text='Synthetic test'):
 """Write a minimal valid multi-page PDF without any third-party PDF dependency."""
 stream=f'BT /F1 12 Tf 20 20 Td ({text}) Tj ET'.encode()
 font=3+2*pages
 objects=[b'<</Type/Catalog/Pages 2 0 R>>',
          f'<</Type/Pages/Kids[{" ".join(f"{3+2*i} 0 R" for i in range(pages))}]/Count {pages}>>'.encode()]
 for i in range(pages):
  objects.append(f'<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]/Resources<</Font<</F1 {font} 0 R>>>>/Contents {4+2*i} 0 R>>'.encode())
  objects.append(b'<</Length '+str(len(stream)).encode()+b'>>stream\n'+stream+b'\nendstream')
 objects.append(b'<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>')
 out=bytearray(b'%PDF-1.4\n');offsets=[]
 for number,body in enumerate(objects,start=1):
  offsets.append(len(out));out+=f'{number} 0 obj'.encode()+body+b'endobj\n'
 xref=len(out)
 out+=f'xref\n0 {len(objects)+1}\n'.encode()+b'0000000000 65535 f \n'
 for offset in offsets:out+=f'{offset:010d} 00000 n \n'.encode()
 out+=f'trailer<</Size {len(objects)+1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n'.encode()
 path.write_bytes(bytes(out))
simple_pdf(out/'31pages.pdf')
